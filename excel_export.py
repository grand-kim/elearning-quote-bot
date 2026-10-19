"""견적 결과를 엑셀 견적서(.xlsx)로 만듭니다."""

from datetime import date, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import pricing as P

_thin = Side(style="thin", color="999999")
_box = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_head = PatternFill("solid", fgColor="DDE6F0")
_won = '#,##0"원"'

# 견적서 상단에 적을 고객 정보 (값이 있는 항목만 표시)
INFO_LABELS = [
    ("client_name", "수신(기관명)"), ("domain", "기관 유형"), ("audience", "교육 대상"),
    ("course_title", "과정명"), ("contact", "담당자"), ("deadline", "희망 납기"),
    ("budget", "예산 범위"), ("reference", "레퍼런스"), ("notes", "상담 메모"),
]


def make_quote_xlsx(quote, info=None):
    info = info or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "견적서"
    for col, w in zip("ABCDE", [26, 20, 12, 14, 16]):
        ws.column_dimensions[col].width = w

    today = date.today()
    ws["A1"] = "이러닝 콘텐츠 제작 견적서"
    ws["A1"].font = Font(size=16, bold=True)
    ws.merge_cells("A1:E1")
    ws["A1"].alignment = Alignment(horizontal="center")

    meta = [("견적일", today.isoformat()),
            ("유효기간", (today + timedelta(days=P.QUOTE_VALID_DAYS)).isoformat() + " 까지"),
            ("발신", P.COMPANY_NAME)]
    meta += [(label, info.get(key)) for key, label in INFO_LABELS if info.get(key)]
    row = 3
    for k, v in meta:
        ws.cell(row, 1, k).font = Font(bold=True)
        ws.cell(row, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    row += 1
    ws.cell(row, 1, "합계 금액 (VAT 포함)").font = Font(bold=True, size=12)
    c = ws.cell(row, 2, quote["합계"])
    c.number_format, c.font = _won, Font(bold=True, size=12)
    row += 2

    # 제작 조건
    ws.cell(row, 1, "■ 제작 조건").font = Font(bold=True)
    row += 1
    for k, v in quote["조건"].items():
        ws.cell(row, 1, k).border = _box
        ws.cell(row, 2, v).border = _box
        row += 1
    row += 1

    # 인건비 내역
    ws.cell(row, 1, "■ 직접인건비 내역").font = Font(bold=True)
    row += 1
    for i, h in enumerate(["항목", "적용 직무", "투입일수", "일 단가", "금액"], 1):
        cell = ws.cell(row, i, h)
        cell.font, cell.fill, cell.border = Font(bold=True), _head, _box
    row += 1
    for line in quote["인건비 내역"]:
        vals = [line["항목"], line["적용 직무"], line["투입일수"], line["일 단가"], line["금액"]]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(row, i, v)
            cell.border = _box
            if i >= 4:
                cell.number_format = _won
        row += 1
    row += 1

    # 직접경비
    if quote["직접경비 내역"]:
        ws.cell(row, 1, "■ 직접경비 내역").font = Font(bold=True)
        row += 1
        for line in quote["직접경비 내역"]:
            ws.cell(row, 1, line["항목"]).border = _box
            ws.cell(row, 2, line["수량"]).border = _box
            cell = ws.cell(row, 5, line["금액"])
            cell.number_format, cell.border = _won, _box
            row += 1
        row += 1

    # 합계
    ws.cell(row, 1, "■ 금액 집계").font = Font(bold=True)
    row += 1
    totals = [
        ("직접인건비", quote["직접인건비"]),
        (f"제경비 (직접인건비 × {P.OVERHEAD_RATE:.0%})", quote["제경비"]),
        (f"기술료 ((인건비+제경비) × {P.TECH_FEE_RATE:.0%})", quote["기술료"]),
        ("직접경비", quote["직접경비"]),
        (f"공급가액 ({P.ROUND_UNIT:,}원 미만 절사)", quote["공급가액"]),
        ("부가세", quote["부가세"]),
        ("합계", quote["합계"]),
    ]
    for k, v in totals:
        ws.cell(row, 1, k).border = _box
        cell = ws.cell(row, 5, v)
        cell.number_format, cell.border = _won, _box
        if k == "합계":
            ws.cell(row, 1).font = cell.font = Font(bold=True)
        row += 1

    row += 1
    ws.cell(row, 1, f"※ 단가 기준: {quote['단가 기준']}")
    ws.cell(row + 1, 1, "※ 본 견적은 상담 내용 기반 예상 금액이며, 최종 금액은 담당자 협의 후 확정됩니다.")
    ws.cell(row + 2, 1, f"※ 문의: {P.COMPANY_NAME} {P.COMPANY_CONTACT}")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
